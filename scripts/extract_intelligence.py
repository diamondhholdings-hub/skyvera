#!/usr/bin/env python3
"""
Extract the operational detail of the budget workbook into
src/data/intelligence-snapshot.json: renewals, the Salesforce renewal
pipeline, receivables, non-recurring revenue, vendors, central-function
charges, headcount, AI cost, product P&Ls, cost benchmarks and the
consolidated income statement history.

The P&L summary and customer list stay in skyvera-snapshot.json
(parse_excel_to_json.py); this file carries everything else.

Headcount is aggregated (by department, location, vertical, team): the
input sheets list named individuals with their pay, which never leaves
this script.

Usage:
    python3 scripts/extract_intelligence.py > src/data/intelligence-snapshot.json
"""

import json
import re
import sys
from collections import defaultdict
from datetime import date, datetime

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

EXCEL_FILE = "2026-09-24 Skyvera Budget P&L - Q4'26 - Final for Todd.xlsx"

QUARTER_RE = re.compile(r"Q[1-4]'\d{2}")
NOTE_LIMIT = 400


def log(message):
    print(message, file=sys.stderr)


def col(letter):
    """Zero-based index of a column letter, for indexing row tuples."""
    return column_index_from_string(letter) - 1


def num(value):
    if isinstance(value, bool):
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    return 0.0


def text(value, limit=None):
    if value is None:
        return None
    s = str(value).strip()
    if not s or s.startswith('#'):  # #REF!, #N/A
        return None
    return s[:limit] if limit else s


def iso(value):
    if isinstance(value, (datetime, date)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    return None


def usd_text(value):
    """'USD 180,936.00' -> 180936.0"""
    if isinstance(value, (int, float)):
        return float(value)
    if not value:
        return 0.0
    cleaned = re.sub(r"[^0-9.\-]", "", str(value))
    try:
        return float(cleaned) if cleaned else 0.0
    except ValueError:
        return 0.0


def rows(ws, first, last=None):
    return ws.iter_rows(min_row=first, max_row=last, values_only=True)


def quarter_labels(ws, row, first_letter, count):
    """Quarter labels from a header row, e.g. ["Q4'26", "Q1'27", ...]."""
    start = column_index_from_string(first_letter)
    labels = []
    for i in range(count):
        v = ws.cell(row, start + i).value
        m = QUARTER_RE.search(str(v)) if v else None
        labels.append(m.group(0) if m else str(v))
    return labels


def round_all(obj):
    if isinstance(obj, float):
        return round(obj, 2)
    if isinstance(obj, dict):
        return {k: round_all(v) for k, v in obj.items()}
    if isinstance(obj, list):
        return [round_all(v) for v in obj]
    return obj


# --------------------------------------------------------------------------
# Revenue: renewals, pipeline, receivables, NRR
# --------------------------------------------------------------------------

def extract_renewals(wb):
    """RR Input: one row per subscription, all business units."""
    ws = wb['RR Input']
    out = []
    for r in rows(ws, 11):
        sub_id = r[col('D')]
        # Adjustment lines below the subscriptions have a note instead of an ID
        if not isinstance(sub_id, (int, float)) or not r[col('B')]:
            continue
        out.append({
            'bu': text(r[col('A')]),
            'customer': text(r[col('B')]),
            'endUser': text(r[col('C')]),
            'subscriptionId': int(sub_id),
            'start': iso(r[col('E')]),
            'end': iso(r[col('F')]),
            'arr': num(r[col('G')]),
            'renewalDate': iso(r[col('H')]),
            'renewalQuarter': text(r[col('I')]),
            'willRenew': text(r[col('J')]),
            'upDownSellPct': num(r[col('K')]) * 100,
            'projectedArr': num(r[col('L')]),
            'arOver90': num(r[col('AG')]),
        })
    log(f"  renewals: {len(out)} subscriptions")
    return out


def extract_pipeline(wb):
    """Salesforce data: renewal opportunities. Rows come in pairs per
    opportunity (the second with a zero price), so keep one per ID."""
    ws = wb['Salesforce data']
    by_id = {}
    for r in rows(ws, 3):
        opp_id = text(r[col('I')])
        account = text(r[col('H')])
        if not opp_id or not account:
            continue
        entry = {
            'opportunityId': opp_id,
            'opportunity': text(r[col('J')]),
            'account': account,
            'bu': text(r[col('C')]),
            'productFamily': text(r[col('N')]),
            'stage': text(r[col('L')]),
            'type': text(r[col('M')]),
            'renewalDate': iso(r[col('K')]),
            'closeDate': iso(r[col('O')]),
            'owner': text(r[col('R')]),
            'probabilityPct': num(r[col('S')]) * 100,  # stored as a fraction
            'currentArr': usd_text(r[col('T')]),
            'arr': usd_text(r[col('U')]),
            'totalPrice': usd_text(r[col('V')]),
            'willRenew': text(r[col('Z')]),
        }
        kept = by_id.get(opp_id)
        if kept is None or (kept['totalPrice'] == 0 and entry['totalPrice'] != 0):
            by_id[opp_id] = entry
    out = list(by_id.values())
    log(f"  pipeline: {len(out)} opportunities")
    return out


def extract_receivables(wb):
    """AR Aging: one row per customer with AR over 90 days (> $5k).
    The sheet carries no aging buckets, only the over-90-day balance."""
    ws = wb['AR Aging']
    out = []
    for r in rows(ws, 4):
        customer = text(r[col('F')])
        if not customer:
            break
        arr = num(r[col('J')])
        out.append({
            'customer': customer,
            'bu': text(r[col('D')]),
            'arOver90': num(r[col('G')]),
            'arrInBudget': arr,
            'arOver10PctOfArr': text(r[col('K')]) == 'Yes',
            'revenueInBudget': (text(r[col('M')]) or '').lower().startswith('include'),
            'latestNote': text(r[col('H')], NOTE_LIMIT),
            'buComments': text(r[col('N')], NOTE_LIMIT),
        })
    out.sort(key=lambda x: x['arOver90'], reverse=True)
    log(f"  receivables: {len(out)} customers, ${sum(x['arOver90'] for x in out):,.0f} over 90 days")
    return out


def extract_nrr(wb):
    """NRR Input: one row per customer per class (consulting, product)."""
    ws = wb['NRR Input']
    quarters = quarter_labels(ws, 5, 'I', 4)
    out = []
    for r in rows(ws, 6):
        name = text(r[col('B')])
        if not name:
            continue
        m = re.search(r'<(.+?)>', name)
        out.append({
            'customer': (m.group(1) if m else name).strip(),
            'class': text(r[col('C')]),
            'department': text(r[col('D')]),
            'nature': text(r[col('E')]),
            'trailing12Months': num(r[col('F')]),
            'byQuarter': {q: num(r[col('I') + i]) for i, q in enumerate(quarters)},
            'notes': text(r[col('N')], NOTE_LIMIT),
        })
    log(f"  nrr: {len(out)} lines")
    return {'quarters': quarters, 'lines': out}


# --------------------------------------------------------------------------
# Costs: vendors, central functions, headcount, AI
# --------------------------------------------------------------------------

def extract_vendors(wb):
    """Vendor Input, using the final budget mapping (columns O-S)."""
    ws = wb['Vendor Input']
    quarters = quarter_labels(ws, 6, 'T', 4)
    out = []
    for r in rows(ws, 7):
        vendor = text(r[col('C')])
        if not vendor:
            continue
        out.append({
            'vendor': vendor,
            'bu': text(r[col('S')]),
            'category': text(r[col('O')]),
            'nature': text(r[col('P')]),
            'department': text(r[col('Q')]),
            'priorQuarterForecast': num(r[col('M')]),
            'priorQuarterBudget': num(r[col('N')]),
            'byQuarter': {q: num(r[col('T') + i]) for i, q in enumerate(quarters)},
            'fourQuarterTotal': num(r[col('AB')]),
            'comments': text(r[col('X')], NOTE_LIMIT),
        })
    out.sort(key=lambda x: x['byQuarter'][quarters[0]], reverse=True)
    log(f"  vendors: {len(out)} lines, ${sum(v['byQuarter'][quarters[0]] for v in out):,.0f} in {quarters[0]}")
    return {'quarters': quarters, 'priorQuarter': quarter_labels(ws, 6, 'M', 1)[0], 'lines': out}


def extract_central_functions(wb):
    """CF: Skyvera central services charged to each vertical.
    Columns M-R are cost; G-L (same headers) are volume."""
    ws = wb['CF']
    labels = [str(ws.cell(7, column_index_from_string('M') + i).value) for i in range(6)]
    out = []
    for r in rows(ws, 8):
        service = text(r[col('F')])
        if not service:
            continue
        out.append({
            'category': text(r[col('A')]),
            'bu': text(r[col('B')]),
            'department': text(r[col('E')]),
            'service': service,
            'group': text(r[col('S')]),
            'cost': {label: num(r[col('M') + i]) for i, label in enumerate(labels)},
        })
    log(f"  central functions: {len(out)} lines")
    return {'periods': labels, 'lines': out}


def extract_headcount(wb):
    """HC Budget Input aggregated; HC Slide for the team view and the
    prior-quarter comparison. No row-level personal data is emitted."""
    ws = wb['HC Budget Input']
    quarters = quarter_labels(ws, 6, 'W', 4)
    current = quarters[0]

    def bucket():
        return {'headcount': 0.0, 'averageHeadcount': 0.0, 'cost': {q: 0.0 for q in quarters}}

    groups = {k: defaultdict(bucket) for k in ('department', 'location', 'vertical', 'costType', 'employment')}
    total = bucket()
    adjustments = []
    terminations = 0

    for r in rows(ws, 7, 200):
        kind = text(r[col('A')])
        name = text(r[col('B')])
        if not name:
            continue
        cost = {q: num(r[col('W') + i]) for i, q in enumerate(quarters)}
        if kind == 'Others':
            # Allocation and fee lines: cost without heads. Describe them by
            # purpose; a hire fee line names the hire, so drop the name.
            label = 'Hire fee' if name.lower().startswith('hire fee') else name
            adjustments.append({'description': label, 'vertical': text(r[col('BC')]), 'cost': cost})
            continue
        if text(r[col('Q')]) == 'Terminate':
            terminations += 1
        keys = {
            'department': text(r[col('F')]) or 'Unassigned',
            'location': text(r[col('J')]) or 'Unknown',
            'vertical': text(r[col('BC')]) or 'Unknown',
            'costType': text(r[col('AH')]) or 'Unknown',
            'employment': 'XO contractor' if kind == 'XO' else 'Non-XO',
        }
        hc, avg = num(r[col('AG')]), num(r[col('AK')])
        for target in [total] + [groups[g][keys[g]] for g in groups]:
            target['headcount'] += hc
            target['averageHeadcount'] += avg
            for q in quarters:
                target['cost'][q] += cost[q]

    # Team view and prior-quarter comparison from HC Slide (already aggregated)
    slide = wb['HC Slide']
    summary = {}
    for r in rows(slide, 4, 8):
        label = text(r[col('C')]) or text(r[col('D')])
        if label:
            summary[label] = {
                'headcountEndOfQuarter': num(r[col('E')]),
                'averageHeadcount': num(r[col('G')]),
                'cost': num(r[col('H')]),
                'priorQuarterHeadcount': num(r[col('I')]),
                'priorQuarterCost': num(r[col('K')]),
            }
    teams = []
    for r in rows(slide, 14, 60):
        if text(r[col('A')]) != 'Yes' or not text(r[col('D')]):
            continue
        teams.append({
            'team': text(r[col('D')]),
            'category': text(r[col('C')]),
            'headcount': num(r[col('E')]),
            'cost': num(r[col('H')]),
            'priorQuarterHeadcount': num(r[col('I')]),
            'priorQuarterCost': num(r[col('K')]),
        })

    def as_list(group):
        return sorted(({'name': k, **v} for k, v in group.items()), key=lambda x: -x['cost'][current])

    log(f"  headcount: {total['headcount']:.0f} at quarter end, ${total['cost'][current]:,.0f} in {current}")
    return {
        'quarters': quarters,
        'total': total,
        'byDepartment': as_list(groups['department']),
        'byLocation': as_list(groups['location']),
        'byVertical': as_list(groups['vertical']),
        'byCostType': as_list(groups['costType']),
        'byEmployment': as_list(groups['employment']),
        'adjustments': adjustments,
        'plannedTerminations': terminations,
        'summary': summary,
        'teams': teams,
    }


def extract_ai_cost(wb):
    """HC + AI cost Metrics: HC and AI tooling cost vs revenue, year over year."""
    ws = wb['HC + AI cost Metrics']
    current, year_ago = str(ws.cell(2, 3).value), str(ws.cell(2, 4).value)
    metrics = {}
    for r in rows(ws, 3, 9):
        label = text(r[col('B')])
        if label and label not in ('Cost',):
            scale = 100 if 'on Revenue' in label else 1  # a ratio: store as a percent
            metrics[label] = {current: num(r[col('C')]) * scale, year_ago: num(r[col('D')]) * scale}
    vendors = []
    for r in rows(ws, 13, 40):
        name = text(r[col('G')])
        if name and (num(r[col('H')]) or num(r[col('I')])):
            vendors.append({'vendor': name, current: num(r[col('H')]), year_ago: num(r[col('I')])})
    return {'current': current, 'yearAgo': year_ago, 'metrics': metrics, 'vendors': vendors}


# --------------------------------------------------------------------------
# Products, benchmarks, margin targets, history
# --------------------------------------------------------------------------

PRODUCT_FIELDS = [
    'product', 'recurringRevenue', 'nonRecurringRevenue', 'totalRevenue',
    'hcCogs', 'allocatedHcCogs', 'totalHcCogs', 'hcOpex', 'allocatedHcOpex', 'totalHcOpex',
    'nhcCogs', 'nhcOpex', 'cfCogs', 'cfOpex', 'writeOff', 'coreAllocation',
    'totalExpenses', 'netProfit', 'netMarginPct',
]


def extract_products(wb):
    """P&Ls by Product: the current and forward quarter budget blocks."""
    ws = wb['P&Ls by Product']
    periods = {}
    for start in ('AK', 'BE', 'BY', 'CS'):
        c = column_index_from_string(start)
        label = str(ws.cell(3, c + 1).value or start)
        lines = []
        for r in rows(ws, 5, 15):
            product = text(r[c - 1])
            if not product:
                continue
            values = [num(r[c - 1 + i]) for i in range(1, len(PRODUCT_FIELDS))]
            line = dict(zip(PRODUCT_FIELDS[1:], values))
            line['netMarginPct'] *= 100
            if line['totalRevenue'] or line['totalExpenses']:
                lines.append({'product': product, **line})
        periods[label] = lines
    log(f"  products: {', '.join(periods)}")
    return periods


def extract_benchmark(wb):
    """Benchmark by Product: each product's cost structure as % of revenue
    against the internal model, and the $ variance to that model."""
    ws = wb['Benchmark by Product']
    products = []
    for c in range(column_index_from_string('C'), column_index_from_string('O') + 1):
        products.append((c, text(ws.cell(1, c).value) or 'Cloudsense-Totogi'))

    def section(first, last):
        out = {}
        for r in range(first, last + 1):
            label = text(ws.cell(r, 2).value)
            if label:
                out[label] = {name: num(ws.cell(r, c).value) * 100 for c, name in products}
        return out

    summary = {}
    for r, key in ((2, 'marginTarget'), (3, 'margin'), (4, 'central'), (14, 'edge'), (26, 'hardCogs'), (29, 'totalCost')):
        summary[key] = {name: num(ws.cell(r, c).value) * 100 for c, name in products}

    variance = {}
    for r in range(72, 86):
        label = text(ws.cell(r, 2).value)
        if label:
            variance[label] = {name: num(ws.cell(r, c).value) for c, name in products[1:]}
    variance['Total'] = {name: num(ws.cell(71, c).value) for c, name in products[1:]}

    rules = []
    slide = wb['Benchmark Slide']
    for r in range(5, 20):
        rule = text(slide.cell(r, column_index_from_string('AH')).value)
        action = text(slide.cell(r, column_index_from_string('AJ')).value) or text(slide.cell(r, column_index_from_string('AI')).value)
        if rule:
            rules.append({'rule': rule, 'action': action})

    return {
        'note': 'Percent of revenue vs the internal cost model (standard benchmark), not peer companies. '
                'Variance is $ per quarter over (+) or under (-) the model.',
        'summaryPct': summary,
        'costPct': section(30, 43),
        'varianceUsd': variance,
        'rules': rules,
    }


def extract_margin_targets(wb, quarter):
    """Margin Target: the current quarter's revenue and target per class (columns H-I)."""
    ws = wb['Margin Target']
    out = []
    for r in rows(ws, 9, 28):
        vertical = text(r[col('A')])
        if not vertical:
            continue
        out.append({
            'vertical': vertical,
            'class': text(r[col('C')]),
            'classType': text(r[col('D')]),
            'revenue': num(r[col('H')]),
            'targetMarginPct': num(r[col('I')]) * 100,
        })
    return {'quarter': quarter, 'lines': out}


def extract_history(wb):
    """Income Statements: consolidated actuals and budget by quarter."""
    ws = wb['Income Statements']
    periods = [(c, text(ws.cell(6, c).value)) for c in range(3, 13)]
    periods = [(c, p) for c, p in periods if p]
    lines = {}
    for r in list(range(9, 32)) + [34]:
        label = text(ws.cell(r, 2).value)
        if not label:
            continue
        is_pct = 'Margin' in label and 'Delta' not in label
        lines[label] = {
            p: (num(ws.cell(r, c).value) * 100 if is_pct else num(ws.cell(r, c).value))
            for c, p in periods
        }
    return {'periods': [p for _, p in periods], 'lines': lines}


def main():
    log(f"Reading {EXCEL_FILE}...")
    wb = load_workbook(EXCEL_FILE, data_only=True)
    snapshot = {
        'source': EXCEL_FILE,
        'renewals': extract_renewals(wb),
        'pipeline': extract_pipeline(wb),
        'receivables': extract_receivables(wb),
        'nrr': extract_nrr(wb),
        'vendors': extract_vendors(wb),
        'centralFunctions': extract_central_functions(wb),
        'headcount': extract_headcount(wb),
        'aiCost': extract_ai_cost(wb),
        'products': extract_products(wb),
        'benchmark': extract_benchmark(wb),
        'marginTargets': None,
        'history': extract_history(wb),
    }
    snapshot['marginTargets'] = extract_margin_targets(wb, snapshot['vendors']['quarters'][0])
    json.dump(round_all(snapshot), sys.stdout, indent=1, default=str)


if __name__ == '__main__':
    main()
